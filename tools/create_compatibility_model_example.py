from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path("outputs/compatibility-model-example/compatibility_model_example.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN = Side(style="thin", color="D9E2F3")


def add_table(ws, name, start_row, headers, rows):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, header)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=start_row + 1):
        for col_index, value in enumerate(row, start=1):
            cell = ws.cell(row_index, col_index, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    end_row = start_row + len(rows)
    end_col = len(headers)
    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    return end_row


def style_sheet(ws, freeze="A2"):
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def fit_columns(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def main():
    wb = Workbook()
    wb.remove(wb.active)

    parts = [
        [1, "PCA-0001", "Maquina de vidro dianteira", "Vidro/Porta", "BMW_X4_X6_MAQUINA_VIDRO"],
        [2, "PCA-0002", "Retrovisor externo", "Retrovisor", "BMW_X4_RETROVISOR"],
    ]

    vehicles = []
    vehicle_id = 1
    for model, years, engines in [
        ("X4", [2009, 2010, 2011, 2012], ["motor 1", "motor 2"]),
        ("X6", [2009, 2010, 2011, 2012], ["motor 1"]),
    ]:
        for year in years:
            for engine in engines:
                vehicles.append([vehicle_id, "BMW", model, year, engine, None, "Aplicacao exata"])
                vehicle_id += 1

    vehicle_lookup = {(row[1], row[2], row[3], row[4]): row[0] for row in vehicles}
    compat = []
    compat_id = 1
    for year in [2009, 2010, 2011, 2012]:
        for engine in ["motor 1", "motor 2"]:
            compat.append([compat_id, 1, vehicle_lookup[("BMW", "X4", year, engine)], 10, "Cadastro em lote: X4 2009-2012 motor 1/2"])
            compat_id += 1
    for year in [2009, 2010, 2011, 2012]:
        compat.append([compat_id, 1, vehicle_lookup[("BMW", "X6", year, "motor 1")], 9, "Cadastro em lote: X6 2009-2012 motor 1"])
        compat_id += 1

    ws = wb.create_sheet("Resumo")
    ws["A1"] = "Modelo recomendado para compatibilidade de pecas"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws["A3"] = "Quantas tabelas principais?"
    ws["B3"] = 3
    ws["A4"] = "Tabelas"
    ws["B4"] = "Parts, Vehicles, PartCompatibilities"
    ws["A5"] = "Ideia"
    ws["B5"] = "A API recebe selecoes agrupadas, expande combinacoes e grava vinculos consultaveis sem busca por nome."
    ws["A7"] = "Por que repetir BMW/X4 por ano/motor?"
    ws["B7"] = "Porque cada linha em Vehicles representa uma aplicacao exata. Isso permite indice, filtro e resposta direta: serve ou nao serve."
    ws["A9"] = "Consulta do anuncio"
    ws["B9"] = "Buscar compatibilidades da peca e ordenar por Relevance DESC."
    for cell in ["A3", "A4", "A5", "A7", "A9"]:
        ws[cell].font = Font(bold=True)
        ws[cell].fill = SUBHEADER_FILL
    ws["B3"].fill = GOOD_FILL
    style_sheet(ws)
    fit_columns(ws, {"A": 28, "B": 95})

    ws = wb.create_sheet("Parts")
    add_table(ws, "PartsTable", 1, ["Id", "OemPartNumber", "Description", "Category", "KeywordGroup"], parts)
    style_sheet(ws)
    fit_columns(ws, {"A": 8, "B": 18, "C": 34, "D": 20, "E": 34})

    ws = wb.create_sheet("Vehicles")
    add_table(ws, "VehiclesTable", 1, ["Id", "Manufacturer", "Model", "ModelYear", "Engine", "Chassis", "Notes"], vehicles)
    style_sheet(ws)
    fit_columns(ws, {"A": 8, "B": 16, "C": 14, "D": 14, "E": 16, "F": 14, "G": 26})

    ws = wb.create_sheet("PartCompatibilities")
    add_table(ws, "CompatTable", 1, ["Id", "PartId", "VehicleId", "Relevance", "Notes"], compat)
    style_sheet(ws)
    fit_columns(ws, {"A": 8, "B": 10, "C": 12, "D": 12, "E": 48})

    ws = wb.create_sheet("API_Input")
    ws["A1"] = "Exemplo de entrada agrupada que a API poderia receber"
    ws["A1"].font = Font(size=14, bold=True, color="1F4E78")
    headers = ["PartId", "Marca", "Modelo", "Anos", "Motores", "Relevance"]
    rows = [
        [1, "BMW", "X4", "2009, 2010, 2011, 2012", "motor 1, motor 2", 10],
        [1, "BMW", "X6", "2009, 2010, 2011, 2012", "motor 1", 9],
    ]
    add_table(ws, "ApiInputTable", 3, headers, rows)
    ws["A8"] = "Regra"
    ws["B8"] = "A API nao grava a lista como texto. Ela expande anos x motores e grava uma linha por veiculo compativel."
    ws["A8"].font = Font(bold=True)
    ws["A8"].fill = NOTE_FILL
    ws["B8"].fill = NOTE_FILL
    style_sheet(ws, freeze="A4")
    fit_columns(ws, {"A": 10, "B": 14, "C": 14, "D": 28, "E": 24, "F": 12})

    ws = wb.create_sheet("Expanded_Result")
    expanded = []
    for row in compat:
        vehicle = next(v for v in vehicles if v[0] == row[2])
        expanded.append([row[1], "PCA-0001", vehicle[1], vehicle[2], vehicle[3], vehicle[4], row[3]])
    expanded.sort(key=lambda r: (-r[6], r[3], r[4], r[5]))
    add_table(ws, "ExpandedResultTable", 1, ["PartId", "Peca", "Marca", "Modelo", "Ano", "Motor", "Relevance"], expanded)
    style_sheet(ws)
    fit_columns(ws, {"A": 10, "B": 16, "C": 12, "D": 12, "E": 10, "F": 16, "G": 12})

    ws = wb.create_sheet("SQL_Model")
    sql_rows = [
        ["Parts", "Id", "PK", "Identificador da peca"],
        ["Parts", "OemPartNumber", "UNIQUE", "Codigo OEM ou codigo interno normalizado"],
        ["Vehicles", "Id", "PK", "Identificador da aplicacao exata"],
        ["Vehicles", "Manufacturer", "INDEX", "Marca, ex: BMW"],
        ["Vehicles", "Model", "INDEX", "Modelo, ex: X4"],
        ["Vehicles", "ModelYear", "INDEX", "Ano do modelo"],
        ["Vehicles", "Engine", "INDEX", "Motor"],
        ["PartCompatibilities", "Id", "PK", "Identificador do vinculo"],
        ["PartCompatibilities", "PartId", "FK + INDEX", "Peca vinculada"],
        ["PartCompatibilities", "VehicleId", "FK + INDEX", "Aplicacao vinculada"],
        ["PartCompatibilities", "Relevance", "CHECK 1..10", "Prioridade da aplicacao no anuncio"],
        ["PartCompatibilities", "PartId + VehicleId", "UNIQUE", "Evita duplicar o mesmo vinculo"],
    ]
    add_table(ws, "SqlModelTable", 1, ["Tabela", "Campo", "Regra", "Descricao"], sql_rows)
    style_sheet(ws)
    fit_columns(ws, {"A": 24, "B": 22, "C": 18, "D": 58})

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
